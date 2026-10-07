from glpi_api import GLPI
from config import CATEGORIA_CONTROLADORIA

from datetime import date, timedelta
import html
import re
import os

from bs4 import BeautifulSoup

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from enviar_email import enviar_relatorio

# ============================================================
# FUNÇÕES
# ============================================================

def calcular_periodo():

    hoje = date.today()

    # Segunda-feira da semana atual
    segunda_atual = hoje - timedelta(days=hoje.weekday())

    # Segunda-feira da semana anterior
    segunda_anterior = segunda_atual - timedelta(days=7)

    # Sexta-feira da semana anterior
    sexta_anterior = segunda_atual - timedelta(days=3)

    data_inicio = f"{segunda_anterior} 00:00:00"
    data_fim = f"{sexta_anterior} 23:59:59"

    return data_inicio, data_fim


def limpar_nome_campo(pergunta):

    pergunta = pergunta.strip()

    # Remove numeração como 1), 2), 10)...
    pergunta = re.sub(r"^\d+\)\s*", "", pergunta)

    # Remove todos os ":" do final
    pergunta = pergunta.rstrip(":").strip()

    return pergunta


def extrair_formulario(content):

    if not content:
        return {}

    # Converte entidades HTML
    content = html.unescape(content)

    soup = BeautifulSoup(
        content,
        "html.parser"
    )

    campos = {}

    for div in soup.find_all("div"):

        b = div.find("b")

        if not b:
            continue

        pergunta = b.get_text(
            " ",
            strip=True
        )

        # Remove a tag <b> para sobrar somente a resposta
        b.extract()

        resposta = div.get_text(
            " ",
            strip=True
        )

        pergunta = limpar_nome_campo(
            pergunta
        )

        resposta = resposta.strip()

        if pergunta:

            campos[pergunta] = resposta

    return campos


def nome_status(status):

    status = str(status)

    status_map = {

        "1": "Novo",
        "2": "Em atendimento",
        "3": "Planejado",
        "4": "Pendente",
        "5": "Solucionado",
        "6": "Fechado"

    }

    return status_map.get(
        status,
        status
    )


# ============================================================
# INÍCIO
# ============================================================

glpi = GLPI()

try:

    print("=" * 70)
    print("CONECTANDO À API DO GLPI")
    print("=" * 70)

    glpi.conectar()


    # ========================================================
    # PERÍODO
    # ========================================================

    data_inicio, data_fim = calcular_periodo()

    print()
    print("=" * 70)
    print("PERÍODO DO RELATÓRIO")
    print("=" * 70)

    print(f"Início: {data_inicio}")
    print(f"Fim:    {data_fim}")


    # ========================================================
    # PESQUISA DOS CHAMADOS
    # ========================================================

    params = {

        # Data de abertura >= início
        "criteria[0][field]": 15,
        "criteria[0][searchtype]": "morethan",
        "criteria[0][value]": data_inicio,


        # Data de abertura <= fim
        "criteria[1][link]": "AND",
        "criteria[1][field]": 15,
        "criteria[1][searchtype]": "lessthan",
        "criteria[1][value]": data_fim,


        # Categoria = Controladoria
        "criteria[2][link]": "AND",
        "criteria[2][field]": 7,
        "criteria[2][searchtype]": "equals",
        "criteria[2][value]": CATEGORIA_CONTROLADORIA,


        "range": "0-999",


        # Campos exibidos na pesquisa
        "forcedisplay[0]": 2,   # ID
        "forcedisplay[1]": 1,   # Título
        "forcedisplay[2]": 15,  # Data de abertura
        "forcedisplay[3]": 7    # Categoria
    }


    resultado = glpi.buscar_chamados(
        params
    )

    chamados = resultado.get(
        "data",
        []
    )


    print()
    print(f"Chamados encontrados: {len(chamados)}")


    # ========================================================
    # PROCESSAR CHAMADOS
    # ========================================================

    dados = []


    for item in chamados:

        chamado_id = item["2"]

        print(
            f"Processando chamado {chamado_id}..."
        )


        chamado = glpi.obter_chamado(
            chamado_id
        )


        campos = extrair_formulario(
            chamado.get(
                "content",
                ""
            )
        )


        linha = {

            "Chamado":
                chamado.get("id"),

            "Título do chamado":
                chamado.get("name"),

            "Status":
                nome_status(
                    chamado.get("status")
                ),

            "Data de abertura":
                chamado.get("date"),

            "Data de fechamento":
                chamado.get("closedate"),


            "Código SAP do fornecedor":
                campos.get(
                    "Código SAP do fornecedor",
                    ""
                ),


            "Tipo de Documento":
                campos.get(
                    "Tipo de Documento",
                    ""
                ),


            "Motivo":
                campos.get(
                    "Motivo",
                    ""
                ),


            "Pedido de Compra":
                campos.get(
                    "Pedido de Compra",
                    ""
                ),


            "Descrição do motivo":
                campos.get(
                    "Descreva o motivo",
                    campos.get(
                        "Descreva o motivo:",
                        ""
                    )
    ),


            "Nº Doc":
                campos.get(
                    "Nº Doc",
                    ""
                ),


            "Já Vencido?":
                campos.get(
                    "Já Vencido?",
                    ""
                ),


            "Empresa":
                campos.get(
                    "Empresa",
                    campos.get("Empresa:", ""
                    )
                ),


            "Observação":
                campos.get(
                    "Observação",
                    campos.get("Observação", ""
                    )
                )
        }


        dados.append(linha)


# ============================================================
# GERAR EXCEL
# ============================================================

    print()
    print("=" * 70)
    print("GERANDO EXCEL")
    print("=" * 70)


    wb = Workbook()


    # ========================================================
    # ABA CHAMADOS
    # ========================================================

    ws = wb.active
    ws.title = "Chamados"


    colunas = [

        "Chamado",
        "Título do chamado",
        "Status",
        "Data de abertura",
        "Data de fechamento",
        "Código SAP do fornecedor",
        "Tipo de Documento",
        "Motivo",
        "Pedido de Compra",
        "Descrição do motivo",
        "Nº Doc",
        "Já Vencido?",
        "Empresa",
        "Observação"

    ]


    # Cabeçalho
    for coluna, nome in enumerate(
        colunas,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=coluna,
            value=nome
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # Dados
    for linha_numero, linha in enumerate(
        dados,
        start=2
    ):

        for coluna, nome in enumerate(
            colunas,
            start=1
        ):

            ws.cell(
                row=linha_numero,
                column=coluna,
                value=linha.get(
                    nome,
                    ""
                )
            )


    # Filtro
    if dados:

        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(colunas))}"
            f"{len(dados) + 1}"
        )


    # Congelar cabeçalho
    ws.freeze_panes = "A2"


    # Largura das colunas
    larguras = {

        "A": 12,
        "B": 35,
        "C": 18,
        "D": 22,
        "E": 22,
        "F": 35,
        "G": 25,
        "H": 20,
        "I": 30,
        "J": 20,
        "K": 40,
        "L": 20,
        "M": 18,
        "N": 18,
        "O": 30,
        "P": 20,
        "Q": 50

    }


    for coluna, largura in larguras.items():

        ws.column_dimensions[
            coluna
        ].width = largura


    # ========================================================
    # ABA RESUMO
    # ========================================================

    resumo = wb.create_sheet(
        "Resumo"
    )


    resumo["A1"] = "RELATÓRIO DE CHAMADOS - CONTROLADORIA"
    resumo["A1"].font = Font(
        bold=True,
        size=14
    )


    resumo["A3"] = "Período"
    resumo["B3"] = (
        f"{data_inicio} até {data_fim}"
    )


    resumo["A4"] = "Total de chamados"
    resumo["B4"] = len(dados)


    # ========================================================
    # CONTADORES
    # ========================================================

    fechados = 0
    abertos = 0
    vencidos = 0
    nao_vencidos = 0


    empresas = {}
    tipos_documento = {}


    for linha in dados:

        status = linha["Status"]

        if status == "Fechado":

            fechados += 1

        else:

            abertos += 1


        vencimento = str(
            linha["Já Vencido?"]
        ).strip().lower()


        if vencimento == "sim":

            vencidos += 1

        elif vencimento == "não" or vencimento == "nao":

            nao_vencidos += 1


        empresa = linha["Empresa"]

        if empresa:

            empresas[empresa] = (
                empresas.get(empresa, 0) + 1
            )


        tipo = linha["Tipo de Documento"]

        if tipo:

            tipos_documento[tipo] = (
                tipos_documento.get(tipo, 0) + 1
            )


    # ========================================================
    # RESUMO GERAL
    # ========================================================

    resumo["A6"] = "Chamados abertos"
    resumo["B6"] = abertos

    resumo["A7"] = "Chamados fechados"
    resumo["B7"] = fechados

    resumo["A8"] = "Já vencidos"
    resumo["B8"] = vencidos

    resumo["A9"] = "Não vencidos"
    resumo["B9"] = nao_vencidos


    # ========================================================
    # POR EMPRESA
    # ========================================================

    linha = 12

    resumo.cell(
        row=linha,
        column=1,
        value="CHAMADOS POR EMPRESA"
    ).font = Font(
        bold=True
    )

    linha += 1


    resumo.cell(
        row=linha,
        column=1,
        value="Empresa"
    ).font = Font(
        bold=True
    )

    resumo.cell(
        row=linha,
        column=2,
        value="Quantidade"
    ).font = Font(
        bold=True
    )

    linha += 1


    for empresa, quantidade in sorted(
        empresas.items()
    ):

        resumo.cell(
            row=linha,
            column=1,
            value=empresa
        )

        resumo.cell(
            row=linha,
            column=2,
            value=quantidade
        )

        linha += 1


    # ========================================================
    # POR TIPO DE DOCUMENTO
    # ========================================================

    linha += 2

    resumo.cell(
        row=linha,
        column=1,
        value="CHAMADOS POR TIPO DE DOCUMENTO"
    ).font = Font(
        bold=True
    )

    linha += 1


    resumo.cell(
        row=linha,
        column=1,
        value="Tipo de Documento"
    ).font = Font(
        bold=True
    )

    resumo.cell(
        row=linha,
        column=2,
        value="Quantidade"
    ).font = Font(
        bold=True
    )

    linha += 1


    for tipo, quantidade in sorted(
        tipos_documento.items()
    ):

        resumo.cell(
            row=linha,
            column=1,
            value=tipo
        )

        resumo.cell(
            row=linha,
            column=2,
            value=quantidade
        )

        linha += 1


    resumo.column_dimensions["A"].width = 35
    resumo.column_dimensions["B"].width = 25


    # ========================================================
    # SALVAR
    # ========================================================

    nome_arquivo = (
        f"Relatorio_Controladoria_"
        f"{data_inicio[:10]}_a_"
        f"{data_fim[:10]}.xlsx"
    )


    wb.save(nome_arquivo)

    caminho_arquivo = os.path.abspath(nome_arquivo)

    enviar_relatorio(
        arquivo=caminho_arquivo,
        destinatarios=[
            "..."
        ],
        assunto="Relatório semanal - Controladoria",
        corpo="""Bom dia,

    Segue em anexo o relatório semanal de chamados da Controladoria.

    Atenciosamente,
    TI
    """
    )


    print()
    print("=" * 70)
    print("RELATÓRIO GERADO COM SUCESSO")
    print("=" * 70)

    print(
        f"Arquivo: {nome_arquivo}"
    )


finally:

    glpi.desconectar()